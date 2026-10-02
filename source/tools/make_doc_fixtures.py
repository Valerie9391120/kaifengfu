# 生成文档功能测试用的样本（都是编的假内容，不含任何私事）：
#   tests/fixtures/文档-测试.md        一份 Markdown
#   tests/fixtures/文档-测试.docx      python-docx 做的 Word：标题、段落、两种列表、表格、一张图
#   tests/fixtures/文档-测试-lo.docx   LibreOffice 存出来的 Word（另一家软件写的，标签的写法不一样）
# 用法（在 source/ 里）：python3 tools/make_doc_fixtures.py
# 要装 python-docx、pillow；第三份还要机器上有 LibreOffice（soffice），没有就跳过。
import io, os, shutil, subprocess, tempfile

from docx import Document
from docx.shared import Inches
from PIL import Image

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "tests", "fixtures")
os.makedirs(OUT, exist_ok=True)

# ---------- Markdown ----------
MD = """# 测试用的采购清单

这是一份给自动测试用的文档，内容是编的。

## 要买的

- 酸奶两盒
- 青菜一把
- **面条**一包

## 备注

1. 先去超市
2. 再去邮局

> 起风的话带伞。
"""
with open(os.path.join(OUT, "文档-测试.md"), "w", encoding="utf-8") as f:
    f.write(MD)

# ---------- Word（python-docx） ----------
doc = Document()
doc.add_heading("测试用的周末计划", level=0)
doc.add_paragraph("这是一份给自动测试用的 Word 文档，内容是编的。第二句用来看一段里有好几句的时候取得全不全。")
doc.add_heading("上午", level=1)
doc.add_paragraph("去早市", style="List Bullet")
doc.add_paragraph("买花", style="List Bullet")
doc.add_heading("下午", level=1)
doc.add_paragraph("洗衣服", style="List Number")
doc.add_paragraph("晒被子", style="List Number")
doc.add_heading("花销", level=2)
t = doc.add_table(rows=3, cols=2)
for r, (a, b) in enumerate([("项目", "钱"), ("花", "12"), ("菜|肉", "30")]):
    t.cell(r, 0).text = a
    t.cell(r, 1).text = b
p = doc.add_paragraph("下面有一张图：")
png = io.BytesIO()
Image.new("RGB", (40, 24), (150, 190, 170)).save(png, format="PNG")
png.seek(0)
doc.add_picture(png, width=Inches(0.8))
p = doc.add_paragraph("最后一段，带")
p.add_run("加粗").bold = True
p.add_run("和 A & B <尖括号> 这样的符号。")
doc.save(os.path.join(OUT, "文档-测试.docx"))

# ---------- Word（LibreOffice 存的） ----------
HTML = """<html><head><meta charset="utf-8"><title>样本</title></head><body>
<h1>另一家软件存的文档</h1>
<p>这一份是 LibreOffice 存出来的，标签的写法和 Word 不一样。</p>
<h2>清单</h2>
<ul><li>第一样</li><li>第二样</li></ul>
<ol><li>先做这个</li><li>再做那个</li></ol>
<table border="1"><tr><td>名字</td><td>数量</td></tr><tr><td>苹果</td><td>3</td></tr></table>
<p>结尾的一段。</p>
</body></html>"""
soffice = shutil.which("soffice") or shutil.which("libreoffice")
if soffice:
    tmp = tempfile.mkdtemp()
    src = os.path.join(tmp, "sample.html")
    with open(src, "w", encoding="utf-8") as f:
        f.write(HTML)
    r = subprocess.run([soffice, "--headless", "--convert-to", "docx:MS Word 2007 XML", "--outdir", tmp, src], capture_output=True, text=True, timeout=180,
                       env=dict(os.environ, HOME=tmp))
    made = os.path.join(tmp, "sample.docx")
    if os.path.exists(made):
        shutil.copy(made, os.path.join(OUT, "文档-测试-lo.docx"))
        print("LibreOffice 那份做好了")
    else:
        print("LibreOffice 没存出来：", (r.stdout + r.stderr)[-300:])
    shutil.rmtree(tmp, ignore_errors=True)
else:
    print("机器上没有 LibreOffice，第三份跳过")

for n in sorted(os.listdir(OUT)):
    print(n, os.path.getsize(os.path.join(OUT, n)))
